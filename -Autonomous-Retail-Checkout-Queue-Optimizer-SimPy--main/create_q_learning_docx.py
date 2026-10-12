import docx
import os

def create_q_learning_doc():
    doc = docx.Document('RL_EXP_3_23BAI1083.docx')

    # Image replacement map (rel_id -> path to new image)
    # Table 6: rId9 -> media/image2.png (CODE)
    # Table 7: rId10 -> media/image3.png (TRAINING PROGRESS)
    # Table 8: rId11 -> media/image4.png (UPDATED MAX)
    # Table 9: rId12 -> media/image5.png (LIVE QUEUE DISPATCH)
    # Table 10: rId13 -> media/image6.png (INITIAL VS FINAL POLICY)
    # Table 11: rId14 -> media/image7.png (TOTAL CUSTOMER OVER TIME)
    
    image_replacements = {
        'rId9': 'q_learning_plots/code_screenshot.png',
        'rId10': 'q_learning_plots/training_progress.png',
        'rId11': 'q_learning_plots/updated_max_q.png',
        'rId12': 'q_learning_plots/live_queue_dispatch.png',
        'rId13': 'q_learning_plots/initial_vs_final_policy.png',
        'rId14': 'q_learning_plots/total_customer_over_time.png',
    }

    # Replace image binary data in document package relationships
    for rel_id, new_img_path in image_replacements.items():
        if rel_id in doc.part.rels:
            target_part = doc.part.rels[rel_id].target_part
            with open(new_img_path, 'rb') as f:
                target_part._blob = f.read()
            print(f"Replaced image for {rel_id} with {new_img_path}")

    # Text updates in paragraphs
    for p in doc.paragraphs:
        full_text = p.text
        if "First-Visit Monte Carlo" in full_text or "Monte Carlo" in full_text:
            # We want to replace Monte Carlo references with Q-Learning appropriately
            pass

    # Let's inspect specific paragraphs by index and update text while preserving runs & styles
    # P35: Ex.No. 3 -> Ex.No. 8 (or Ex.No. 3 Q-Learning)
    p35 = doc.paragraphs[35]
    if "Ex.No. 3" in p35.text:
        p35.runs[0].text = "Ex.No. 8"

    # P36: Title
    p36 = doc.paragraphs[36]
    for r in p36.runs:
        if "On-Policy First-Visit Monte Carlo Control" in r.text:
            r.text = r.text.replace("On-Policy First-Visit Monte Carlo Control", "Off-Policy Tabular Q-Learning Control")
        elif "Monte Carlo" in r.text:
            r.text = r.text.replace("Monte Carlo", "Q-Learning")

    # P37: Aim
    p37 = doc.paragraphs[37]
    for r in p37.runs:
        if "First-Visit Monte Carlo control" in r.text:
            r.text = r.text.replace("First-Visit Monte Carlo control", "1-Step Tabular Q-Learning control")
        elif "Monte Carlo" in r.text:
            r.text = r.text.replace("Monte Carlo", "Q-Learning")

    # P38: Problem Statement
    p38 = doc.paragraphs[38]
    for r in p38.runs:
        if "episodic return backtracking" in r.text:
            r.text = r.text.replace("episodic return backtracking", "temporal difference error updates")
        elif "sample returns" in r.text:
            r.text = r.text.replace("sample returns", "1-step temporal difference updates")

    # P39: Implementation
    p39 = doc.paragraphs[39]
    for r in p39.runs:
        if "On-Policy First-Visit Monte Carlo Control" in r.text:
            r.text = r.text.replace("On-Policy First-Visit Monte Carlo Control", "Off-Policy Tabular Q-Learning Control")
        elif "incremental sample-mean updates" in r.text:
            r.text = r.text.replace("incremental sample-mean updates", "1-step temporal difference (TD) updates")

    # P42: Objectives
    p42 = doc.paragraphs[42]
    for r in p42.runs:
        if "episodic return backtracking (First-Visit Monte Carlo updates)" in r.text:
            r.text = r.text.replace("episodic return backtracking (First-Visit Monte Carlo updates)", "1-step Temporal Difference (Q-Learning) updates")
        elif "First-Visit Monte Carlo" in r.text:
            r.text = r.text.replace("First-Visit Monte Carlo", "1-step Q-Learning")

    # P53: Action-Value Function description
    p53 = doc.paragraphs[53]
    for r in p53.runs:
        if "at the end of each trajectory via First-Visit Monte Carlo incremental mean updates" in r.text:
            r.text = r.text.replace("at the end of each trajectory via First-Visit Monte Carlo incremental mean updates", "at each transition step via 1-Step Temporal Difference Q-Learning updates")
        elif "First-Visit Monte Carlo" in r.text:
            r.text = r.text.replace("First-Visit Monte Carlo", "1-Step Q-Learning")

    # P55: Optimal Policy
    p55 = doc.paragraphs[55]
    for r in p55.runs:
        if "First-Visit Monte Carlo" in r.text:
            r.text = r.text.replace("First-Visit Monte Carlo", "Q-Learning")

    # P62: Heading above code block 2
    p62 = doc.paragraphs[62]
    for r in p62.runs:
        if "Monte Carlo Queue Optimizer Agent" in r.text:
            r.text = r.text.replace("Monte Carlo Queue Optimizer Agent", "Q-Learning Queue Optimizer Agent")

    # P86: Result
    p86 = doc.paragraphs[86]
    # Update text and numbers:
    # "The experiment demonstrates a significant improvement in queue management using the Monte Carlo Reinforcement Learning optimizer compared to random routing. The Monte Carlo RL agent reduced the total customer wait-units from 563 (Baseline Random) to 190, achieving a remarkable 66.25% reduction in queue congestion over 100 simulation steps. This indicates that the optimized policy effectively minimizes waiting times and balances queue loads across checkout counters."
    p86.runs[0].text = "The experiment demonstrates a significant improvement in queue management using the Tabular Q-Learning Reinforcement Learning optimizer compared to random routing. The Q-Learning RL agent reduced the total customer wait-units from "
    p86.runs[1].text = "381 (Baseline Random)"
    p86.runs[2].text = " to "
    p86.runs[3].text = "268"
    p86.runs[4].text = ", achieving a remarkable "
    p86.runs[5].text = "29.66% reduction in queue congestion"
    p86.runs[6].text = " over 100 simulation steps. This indicates that the optimized policy effectively minimizes waiting times and balances queue loads across checkout counters."

    # Update Code Tables (Table 1, 2, 3, 4, 5)
    # Table 1: Environment code (keep RetailQueueEnv)
    # Table 2: Agent code (QLearningQueueOptimizer)
    t2_code = """class QLearningQueueOptimizer:
    def __init__(self, num_counters=3, max_queue_len=6, gamma=0.9, alpha=0.1, epsilon=0.3, epsilon_decay=0.995, min_epsilon=0.05):
        self.num_counters = num_counters
        self.max_queue_len = max_queue_len
        self.gamma = gamma
        self.alpha = alpha
        self.epsilon = epsilon
        self.epsilon_decay = epsilon_decay
        self.min_epsilon = min_epsilon

        # State space dimensions: (max_queue_len + 1) for each queue
        state_dim = tuple([self.max_queue_len + 1] * self.num_counters)

        # Q-table: shape (len_q0, len_q1, len_q2, actions)
        self.Q = np.zeros(state_dim + (self.num_counters,))

        # Policy initialization (deterministic best action, starts random)
        self.policy = np.random.choice(self.num_counters, size=state_dim)

    def choose_action(self, state, explore=True):
        if explore and np.random.rand() < self.epsilon:
            return np.random.choice(self.num_counters)
        return np.argmax(self.Q[state])

    def update_Q(self, state, action, reward, next_state):
        \"\"\"1-Step Temporal Difference Q-Learning Update\"\"\"
        best_next_action = np.argmax(self.Q[next_state])
        td_target = reward + self.gamma * self.Q[next_state][best_next_action]
        td_error = td_target - self.Q[state][action]
        
        # Q-value update rule
        self.Q[state][action] += self.alpha * td_error

        # Greedy policy update
        self.policy[state] = np.argmax(self.Q[state])

    def decay_epsilon(self):
        # Decay exploration rate
        self.epsilon = max(self.min_epsilon, self.epsilon * self.epsilon_decay)"""
    
    doc.tables[2].cell(0, 0).paragraphs[0].text = t2_code

    # Table 3: Training loop code
    t3_code = """env = RetailQueueEnv(num_counters=3, max_queue_len=5, service_rates=[0.6, 0.4, 0.3])
agent = QLearningQueueOptimizer(num_counters=3, max_queue_len=5, gamma=0.95, alpha=0.1, epsilon=0.4)

# Record initial policy snapshot for faculty comparison
initial_policy_snapshot = np.copy(agent.policy)

num_episodes = 250
episode_rewards = []
display_interval = 25

print("Starting Q-Learning Optimization Training...")
time.sleep(1)

for ep in range(1, num_episodes + 1):
    state = env.reset()
    total_ep_reward = 0
    ep_history = []
    
    for step in range(25):
        action = agent.choose_action(state, explore=True)
        next_state, reward = env.step(action)
        
        # 1-Step Q-Learning Update
        agent.update_Q(state, action, reward, next_state)
        
        total_ep_reward += reward
        ep_history.append((state, action, reward))
        state = next_state
        
    agent.decay_epsilon()
    episode_rewards.append(total_ep_reward)

    # Real-time Visual Update
    if ep % display_interval == 0 or ep == 1:
        clear_output(wait=True)
        fig, axes = plt.subplots(1, 3, figsize=(18, 5))

        # 1. Live Learning Curve
        axes[0].plot(range(1, len(episode_rewards) + 1), episode_rewards, color='royalblue', alpha=0.4, label='Raw Return')

        if len(episode_rewards) >= 10:
            smooth_rewards = np.convolve(episode_rewards, np.ones(10)/10, mode='valid')
            axes[0].plot(range(10, len(episode_rewards) + 1), smooth_rewards, color='firebrick', lw=2, label='10-Ep Moving Avg')

        axes[0].set_title(f"Training Progress (Episode {ep}/{num_episodes})", fontweight='bold')
        axes[0].set_xlabel("Episode")
        axes[0].set_ylabel("Total Cumulative Reward")
        axes[0].grid(True, linestyle='--', alpha=0.6)
        axes[0].legend()

        # 2. Updated Q-Value Slice (Counter 2 queue fixed at 0)
        q_slice = np.max(agent.Q[:, :, 0, :], axis=-1)
        sns.heatmap(q_slice, annot=True, fmt=".1f", cmap="magma", ax=axes[1], cbar=True)
        axes[1].set_title(r"Updated Max $Q(s, a)$ Slice (Lane 2 = 0)", fontweight='bold')
        axes[1].set_xlabel("Counter 1 Queue Length")
        axes[1].set_ylabel("Counter 0 Queue Length")

        # 3. Dynamic Queue Simulation Visual
        current_state = ep_history[-1][0]
        chosen_action = ep_history[-1][1]
        bars = axes[2].bar(["Lane 0\\n(Fast: $\\mu$=0.6)", "Lane 1\\n(Mid: $\\mu$=0.4)", "Lane 2\\n(Slow: $\\mu$=0.3)"],
                           current_state, color=['#2ca02c', '#1f77b4', '#ff7f0e'], edgecolor='black', width=0.5)
        axes[2].set_ylim(0, env.max_queue_len + 1)
        axes[2].set_ylabel("Active Customers")
        axes[2].set_title(f"Live Queue Dispatch (Action: Routed to Lane {chosen_action})", fontweight='bold')
        axes[2].grid(axis='y', linestyle='--', alpha=0.7)

        plt.tight_layout()
        plt.show()
        time.sleep(0.05)

print("\\nTraining Complete!")"""

    doc.tables[3].cell(0, 0).paragraphs[0].text = t3_code

    # Table 4: Policy comparison title in heatmap
    t4_code = """fig, axes = plt.subplots(1, 2, figsize=(14, 6))

# Slice index: Counter 2 fixed at 1 customer
fixed_q2 = 1

# Initial Policy Slice
sns.heatmap(initial_policy_snapshot[:, :, fixed_q2], annot=True, cmap="Set3", cbar=False, ax=axes[0],
            linewidths=1, linecolor='gray')
axes[0].set_title(f"Initial (Untrained / Random) Policy\\n[Fixed Lane 2 = {fixed_q2}]", fontsize=12, fontweight='bold')
axes[0].set_xlabel("Counter 1 Queue Length")
axes[0].set_ylabel("Counter 0 Queue Length")

# Final Policy Slice
sns.heatmap(agent.policy[:, :, fixed_q2], annot=True, cmap="Accent", cbar=False, ax=axes[1],
            linewidths=1, linecolor='gray')
axes[1].set_title(f"Final Optimized Policy (Q-Learning)\\n[Fixed Lane 2 = {fixed_q2}]", fontsize=12, fontweight='bold')
axes[1].set_xlabel("Counter 1 Queue Length")
axes[1].set_ylabel("Counter 0 Queue Length")

plt.suptitle("Policy Decision Comparison (Values represent Target Lane: 0, 1, or 2)", fontsize=14, y=1.02)
plt.tight_layout()
plt.show()"""

    doc.tables[4].cell(0, 0).paragraphs[0].text = t4_code

    # Table 5: Benchmark code text update
    t5_code = """def evaluate_policy(env, agent=None, steps=100):
    env.reset()
    total_wait_time = 0
    history = []

    for _ in range(steps):
        state = tuple(env.queues)
        if agent is None:
            action = np.random.choice(env.num_counters)  # Baseline: Random allocation
        else:
            action = np.argmax(agent.Q[state])          # RL Optimal Policy

        next_state, reward = env.step(action)
        total_wait_time += np.sum(next_state)
        history.append(list(next_state))

    return total_wait_time, np.array(history)

# Run benchmark
rand_wait, rand_hist = evaluate_policy(env, agent=None, steps=100)
opt_wait, opt_hist = evaluate_policy(env, agent=agent, steps=100)

print(f"{'='*45}")
print(f"BENCHMARK RESULTS OVER 100 TIMESTEPS:")
print(f"{'='*45}")
print(f"Total Customer Wait-Units (Baseline Random): {rand_wait}")
print(f"Total Customer Wait-Units (Q-Learning RL):  {opt_wait}")
print(f"Queue Congestion Reduction:                 {((rand_wait - opt_wait)/rand_wait)*100:.2f}%")
print(f"{'='*45}")

# Plotting lane occupancies over time
plt.figure(figsize=(14, 5))
plt.plot(np.sum(rand_hist, axis=1), label='Random Baseline Queue Load', color='red', linestyle='--', alpha=0.7)
plt.plot(np.sum(opt_hist, axis=1), label='Q-Learning Optimized Queue Load', color='green', lw=2)
plt.title("Total Customers in Store Over Time", fontsize=13, fontweight='bold')
plt.xlabel("Simulation Steps")
plt.ylabel("Total Queue Length")
plt.legend()
plt.grid(True, linestyle='--', alpha=0.5)
plt.show()"""

    doc.tables[5].cell(0, 0).paragraphs[0].text = t5_code

    # Output file paths
    out_paths = [
        'RL_EXP_8_23BAI1083.docx',
        'RL_EXP_Q_Learning_23BAI1083.docx',
        'RL_EXP_3_23BAI1083_Q_Learning.docx'
    ]

    for p in out_paths:
        doc.save(p)
        print(f"Saved modified document to {p}")

if __name__ == '__main__':
    create_q_learning_doc()
